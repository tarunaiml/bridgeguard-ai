import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def set_footer(slide):
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(9.0), Inches(0.5))
    tf = txBox.text_frame
    p = tf.add_paragraph()
    p.text = "BRIDGEGUARD ML | IDP Review II | Project ID: 202600749"
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(100, 116, 139)

def add_title(slide, title_text):
    tbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.33), Inches(1))
    p = tbox.text_frame.add_paragraph()
    p.text = title_text
    p.font.bold = True
    p.font.size = Pt(36)
    p.font.color.rgb = RGBColor(17, 24, 39)
    return tbox

def add_speaker_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# ---------------------------------------------------------
# SLIDE 1 — TITLE
# ---------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)

title_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11.33), Inches(1.5))
tf = title_box.text_frame
p = tf.add_paragraph()
p.text = "BridgeGuard ML"
p.font.bold = True
p.font.size = Pt(54)
p.font.color.rgb = RGBColor(17, 24, 39)
p.alignment = PP_ALIGN.CENTER

sub_box = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.33), Inches(1))
tf2 = sub_box.text_frame
p2 = tf2.add_paragraph()
p2.text = "Explainable Machine Learning for Bridge Structural Health Monitoring"
p2.font.size = Pt(28)
p2.font.color.rgb = RGBColor(37, 99, 235)
p2.alignment = PP_ALIGN.CENTER

info_box = slide.shapes.add_textbox(Inches(1.5), Inches(3.8), Inches(5), Inches(3))
tf3 = info_box.text_frame
tf3.text = "IDP Review II\nAY 2026–2027\nProject ID: 202600749\n\nInstitution:\nVIT Chennai"
for p in tf3.paragraphs:
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(17, 24, 39)

team_box = slide.shapes.add_textbox(Inches(7.5), Inches(3.8), Inches(5), Inches(3))
tf4 = team_box.text_frame
tf4.text = "Team:\n25BAI1241 – Avula Tarun\n25BAI1139 – N Anudeep\n25BAI1392 – Ayush Sharma\n\nGuide:\nPadma J\nEmployee ID: 53408"
for p in tf4.paragraphs:
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(17, 24, 39)

add_speaker_notes(slide, "Welcome to our IDP Review II presentation. Our project is BridgeGuard ML, which focuses on applying explainable machine learning to bridge structural health monitoring. We are focusing on health-relative feature extraction rather than arbitrary absolute thresholds.")

# ---------------------------------------------------------
# SLIDE 2 — AGENDA
# ---------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)
set_footer(slide)
add_title(slide, "AGENDA")

agenda_items = [
    "1. Problem & Motivation",
    "2. Literature Survey",
    "3. Research Gap",
    "4. Objectives",
    "5. Proposed Methodology",
    "6. ML Algorithm Selection",
    "7. System Architecture",
    "8. Initial Prototype",
    "9. Explainable AI",
    "10. Future Work"
]

left_box = slide.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(5), Inches(5))
tf_left = left_box.text_frame
for item in agenda_items[:5]:
    p = tf_left.add_paragraph()
    p.text = item
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(37, 99, 235)
    p.space_after = Pt(20)

right_box = slide.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5), Inches(5))
tf_right = right_box.text_frame
for item in agenda_items[5:]:
    p = tf_right.add_paragraph()
    p.text = item
    p.font.size = Pt(24)
    p.font.color.rgb = RGBColor(37, 99, 235)
    p.space_after = Pt(20)

add_speaker_notes(slide, "Here is the agenda for our presentation. We will start with the problem motivation, review the literature and research gaps, introduce our proposed methodology and architecture, discuss our ML selection and explainability strategy, and conclude with our initial prototype and future work.")

# ---------------------------------------------------------
# SLIDE 3 — PROBLEM & MOTIVATION
# ---------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)
set_footer(slide)
add_title(slide, "PROBLEM & MOTIVATION")

bbox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.33), Inches(4))
tf = bbox.text_frame
bullets = [
    "Bridges experience continuous dynamic and environmental loading.",
    "Traditional inspection is periodic; sensor-based SHM enables continuous monitoring.",
    "Raw sensor measurements alone are difficult to interpret.",
    "Different bridges have different normal response ranges (geometry, materials, environment).",
    "Fixed absolute thresholds can therefore generate misleading decisions.",
    "Machine Learning can learn patterns from health-relevant features.",
    "Explainable AI can make model predictions easier to interpret for engineers."
]
for b in bullets:
    p = tf.add_paragraph()
    p.text = "• " + b
    p.font.size = Pt(20)
    p.font.color.rgb = RGBColor(17, 24, 39)
    p.space_after = Pt(10)

rc_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12.33), Inches(1.5))
rc_box.fill.solid()
rc_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
rc_box.line.color.rgb = RGBColor(37, 99, 235)
tf = rc_box.text_frame
p = tf.add_paragraph()
p.text = "RESEARCH CHALLENGE"
p.font.bold = True
p.font.size = Pt(16)
p.font.color.rgb = RGBColor(37, 99, 235)
p2 = tf.add_paragraph()
p2.text = "How can ML-based bridge condition detection remain relevant to bridge-specific behavior while providing interpretable predictions?"
p2.font.size = Pt(20)
p2.font.bold = True
p2.font.color.rgb = RGBColor(17, 24, 39)

add_speaker_notes(slide, "Bridges undergo continuous dynamic loading, making periodic manual inspections insufficient. While sensors provide continuous data, raw signals are hard to interpret. Crucially, because every bridge is different, a fixed absolute threshold for safety doesn't work. Our research challenge is determining how to make ML models adaptable to bridge-specific behavior while keeping predictions interpretable.")

# ---------------------------------------------------------
# SLIDES 4–8 — LITERATURE SURVEY (5 slides, 3 papers each)
# ---------------------------------------------------------
lit_data = [
    [
        ("Worden & Manson (2007)", "ML application in SHM", "Pattern recognition is fundamental to SHM.", "Relies heavily on absolute labeled data.", "Use bridge-specific normalization."),
        ("Farrar & Worden (2012)", "SHM machine learning framework", "Damage detection is statistical pattern recognition.", "Environmental variability complicates baseline.", "Focus on health-relative features."),
        ("Avci et al. (2021)", "Vibration-based damage detection", "Deep learning removes manual extraction.", "Lack of explainability in deep models.", "Integrate SHAP for explainability.")
    ],
    [
        ("Bao et al. (2019)", "Computer vision & ML in SHM", "Data-driven methods are highly viable.", "Susceptible to environmental noise.", "Normalize features against baseline."),
        ("Azimi et al. (2020)", "Data-driven SHM via DL", "Neural networks are highly effective.", "Cross-structure generalization is poor.", "Investigate relative features."),
        ("Entezami et al. (2022)", "Unsupervised ML & env variability", "Unsupervised methods handle varying environments.", "Hard to distinguish damage from rare normal events.", "Compare multiple ML algorithms.")
    ],
    [
        ("Neves et al. (2017)", "Model-free anomaly detection", "Avoids complex finite element models.", "Hard to localize damage specifically.", "Use explainable AI to highlight features."),
        ("Malekzadeh et al. (2015)", "ML classification for SHM", "SVM and RF perform well on extracted features.", "Thresholds are strictly bridge-specific.", "Extract health-relative normalized features."),
        ("Sun et al. (2020)", "Deep learning for SHM review", "Autoencoders excel at anomaly detection.", "High computational cost, 'black-box' nature.", "Use lightweight interpretable ML models.")
    ],
    [
        ("Figueiredo et al. (2011)", "Damage detection under variability", "PCA mitigates operational variability.", "Linear assumptions limit some models.", "Compare non-linear models (e.g., Random Forest)."),
        ("Oh et al. (2020)", "Explainable ML approach for SHM", "SHAP values improve trust in models.", "SHAP computation can be expensive.", "Optimize feature set prior to SHAP."),
        ("Toh & Park (2020)", "Review of ML in bridge SHM", "Sensor fusion improves accuracy.", "Limited availability of damaged state data.", "Train on normal baseline deviations.")
    ],
    [
        ("Vamvoudakis-Stefanou (2023)", "Unsupervised damage detection", "Transfer learning helps across populations.", "Requires high structural similarity.", "Analyze generalization limitations."),
        ("Sony et al. (2021)", "Smart sensing technology in SHM", "IoT enables real-time continuous SHM.", "Data synchronization and power constraints.", "Design edge-compatible feature extraction."),
        ("Huang et al. (2019)", "Anomaly detection in bridge data", "Ensemble models give robust anomaly detection.", "Imbalanced data affects model reliability.", "Use anomaly/one-class formulations if needed.")
    ]
]

for i, papers in enumerate(lit_data):
    slide = prs.slides.add_slide(blank_layout)
    set_footer(slide)
    add_title(slide, f"LITERATURE SURVEY ({i+1}/5)")
    
    # Create a table for the 3 papers
    table_box = slide.shapes.add_table(4, 5, Inches(0.5), Inches(1.5), Inches(12.33), Inches(5))
    table = table_box.table
    
    headers = ["Paper / Authors", "Focus", "Key Takeaway", "Limitation / Future Scope", "Our Proposed Response"]
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = h
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(14)
        cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(37, 99, 235)
        
    for row_idx, paper in enumerate(papers):
        for col_idx, val in enumerate(paper):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = val
            cell.text_frame.paragraphs[0].font.size = Pt(12)
            cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(17, 24, 39)
            
    add_speaker_notes(slide, "Our literature survey identifies that while ML is highly effective for SHM, many models treat absolute sensor data as universal, suffer from black-box unexplainability, and struggle to generalize due to environmental variations. Our response emphasizes normalized features and SHAP explainability.")

# ---------------------------------------------------------
# SLIDE 9 — RESEARCH GAP
# ---------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)
set_footer(slide)
add_title(slide, "RESEARCH GAP & PROPOSED DIRECTION")

flow_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.33), Inches(1.0))
tf = flow_box.text_frame
tf.text = "Existing SHM   →   Sensor Data   →   Feature Extraction   →   ML Prediction"
for p in tf.paragraphs:
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(100, 116, 139)
    p.alignment = PP_ALIGN.CENTER

bbox = slide.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(12.33), Inches(2.5))
tf = bbox.text_frame
p = tf.add_paragraph()
p.text = "Identified Limitations:"
p.font.bold = True
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(220, 38, 38)

bullets = [
    "Absolute sensor values are not directly transferable across different bridges.",
    "Environmental and operational variability can heavily affect sensor response.",
    "Fixed thresholds become overly simplistic decision rules.",
    "Some ML predictions lack interpretability (black-box models).",
    "Individual studies often focus solely on one dataset or structure."
]
for b in bullets:
    p = tf.add_paragraph()
    p.text = "• " + b
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(17, 24, 39)

dir_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.5), Inches(12.33), Inches(2.0))
dir_box.fill.solid()
dir_box.fill.fore_color.rgb = RGBColor(239, 246, 255)
tf = dir_box.text_frame
p = tf.add_paragraph()
p.text = "OUR DIRECTION"
p.font.bold = True
p.font.size = Pt(20)
p.font.color.rgb = RGBColor(37, 99, 235)
p.alignment = PP_ALIGN.CENTER

p2 = tf.add_paragraph()
p2.text = "Health-Relative Features  +  Bridge-Specific Baseline  +  ML Algorithm Comparison  +  SHAP Explainability"
p2.font.size = Pt(18)
p2.font.bold = True
p2.font.color.rgb = RGBColor(17, 24, 39)
p2.alignment = PP_ALIGN.CENTER
p2.space_before = Pt(10)

add_speaker_notes(slide, "The literature highlights several gaps: absolute sensor values don't transfer across bridges, and complex ML models are often black boxes. We propose a direction focusing on health-relative features normalized against a bridge-specific baseline, combined with rigorous ML comparison and SHAP explainability.")

# ---------------------------------------------------------
# SLIDE 10 — PROJECT OBJECTIVES
# ---------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)
set_footer(slide)
add_title(slide, "PROJECT OBJECTIVES")

objs = [
    "1. Develop health-relative feature extraction from bridge response data.",
    "2. Establish bridge-specific healthy baselines and normalized health indicators.",
    "3. Compare multiple ML algorithms for structural condition / anomaly detection.",
    "4. Select an appropriate model based on validation performance and suitability.",
    "5. Integrate SHAP-based explainability for ML predictions.",
    "6. Develop a dashboard that communicates prediction, trends and feature contributions."
]

bbox = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.33), Inches(5))
tf = bbox.text_frame
for ob in objs:
    p = tf.add_paragraph()
    p.text = ob
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(17, 24, 39)
    p.space_after = Pt(20)

add_speaker_notes(slide, "Our objectives are highly focused on this gap. We aim to develop health-relative feature extraction, establish healthy baselines, compare various ML algorithms rather than assuming one is best, and integrate SHAP explainability into a working dashboard prototype.")

# ---------------------------------------------------------
# SLIDE 11 — PROPOSED METHODOLOGY
# ---------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)
set_footer(slide)
add_title(slide, "PROPOSED METHODOLOGY")

mbox = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12.33), Inches(5))
tf = mbox.text_frame
flow = [
    "DATA SOURCES (Real SHM datasets & Future physical miniature bridge)",
    "↓",
    "Preprocessing",
    "↓",
    "Health-Relative Feature Extraction (RMS, Peak, StdDev, Freq, Tilt, Temp)",
    "↓",
    "Healthy Baseline / Bridge-Specific Normalization",
    "↓",
    "ML Algorithm Comparison",
    "↓",
    "Selected ML Model",
    "↓",
    "Prediction",
    "↓",
    "SHAP Explainability",
    "↓",
    "Dashboard"
]
for item in flow:
    p = tf.add_paragraph()
    p.text = item
    p.font.size = Pt(14)
    p.font.bold = (item != "↓")
    p.font.color.rgb = RGBColor(37, 99, 235) if item == "↓" else RGBColor(17, 24, 39)
    p.alignment = PP_ALIGN.CENTER

add_speaker_notes(slide, "Our methodology moves from data sources—including real datasets and future miniature models—through preprocessing and health-relative feature extraction. Critically, we perform bridge-specific normalization before comparing ML algorithms. The selected model's predictions will then be explained using SHAP.")

# ---------------------------------------------------------
# SLIDE 12 — SYSTEM ARCHITECTURE
# ---------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)
set_footer(slide)
add_title(slide, "SYSTEM ARCHITECTURE")

arch_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12.33), Inches(5.5))
tf = arch_box.text_frame
arch_flow = [
    "BRIDGE / SENSOR DATA",
    "↓",
    "DATASET / ESP32 SENSOR INPUT (Planned: MPU6050, Temp/Humidity)",
    "↓",
    "PREPROCESSING",
    "↓",
    "HEALTH-RELATIVE FEATURE EXTRACTION",
    "↓",
    "BRIDGE-SPECIFIC NORMALIZATION",
    "↓",
    "ML MODEL COMPARISON",
    "↓",
    "SELECTED ML MODEL",
    "↓",
    "PREDICTION",
    "↓",
    "SHAP EXPLANATION",
    "↓",
    "DASHBOARD"
]
for item in arch_flow:
    p = tf.add_paragraph()
    p.text = item
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(37, 99, 235) if item == "↓" else RGBColor(17, 24, 39)
    p.alignment = PP_ALIGN.CENTER

p_note = tf.add_paragraph()
p_note.text = "\n*Note: Hardware (ESP32/Sensors) is planned / under development. Data transmission flows to ML processing and dashboard."
p_note.font.size = Pt(12)
p_note.font.bold = False
p_note.font.color.rgb = RGBColor(100, 116, 139)
p_note.alignment = PP_ALIGN.CENTER

add_speaker_notes(slide, "This is our proposed system architecture. It outlines the data flow from future physical sensors or current datasets, through our software pipeline of extraction, normalization, and ML processing, ultimately feeding the explanation and prediction into the dashboard.")

# ---------------------------------------------------------
# SLIDE 13 — ML ALGORITHM SELECTION
# ---------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)
set_footer(slide)
add_title(slide, "COMPARATIVE ML MODEL SELECTION")

t_box = slide.shapes.add_table(8, 4, Inches(0.5), Inches(1.5), Inches(12.33), Inches(3.5))
table = t_box.table
headers = ["Algorithm", "Strength", "Limitation", "Role in our study"]
for col_idx, h in enumerate(headers):
    cell = table.cell(0, col_idx)
    cell.text = h
    cell.text_frame.paragraphs[0].font.bold = True
    cell.text_frame.paragraphs[0].font.size = Pt(14)
    cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor(37, 99, 235)

algos = [
    ("Logistic Regression", "Simple, interpretable baseline", "Struggles with non-linear relationships", "Baseline model"),
    ("KNN", "No strict mathematical assumptions", "Computationally heavy at scale", "Distance-based comparison"),
    ("SVM", "Effective in high-dimensional spaces", "Complex hyperparameter tuning", "Margin-based comparison"),
    ("Decision Tree", "Highly interpretable logic", "Prone to overfitting", "Baseline tree model"),
    ("Random Forest", "Robust ensemble, handles non-linearity", "Less interpretable than single tree", "Primary ensemble candidate"),
    ("XGBoost", "High accuracy and performance", "Can overfit on small datasets", "Advanced ensemble candidate"),
    ("Isolation Forest", "Excellent for novelty detection", "Requires clean training baseline", "Anomaly detection candidate")
]
for row_idx, row in enumerate(algos):
    for col_idx, val in enumerate(row):
        cell = table.cell(row_idx + 1, col_idx)
        cell.text = val
        cell.text_frame.paragraphs[0].font.size = Pt(12)

flow_b = slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12.33), Inches(1.5))
tf = flow_b.text_frame
p = tf.add_paragraph()
p.text = "Same feature set  →  Train / Validate  →  Compare metrics  →  Select suitable model\n*Model selection will be strictly based on experimental validation."
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = RGBColor(17, 24, 39)
p.alignment = PP_ALIGN.CENTER

add_speaker_notes(slide, "We are not assuming Random Forest is the best model blindly. We will evaluate multiple candidates including Logistic Regression, SVM, and ensembles like XGBoost and Isolation Forest on the same feature set, and strictly select the final model based on experimental validation metrics.")

# ---------------------------------------------------------
# SLIDE 14 — INITIAL IMPLEMENTATION / PROTOTYPE
# ---------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)
set_footer(slide)
add_title(slide, "INITIAL IMPLEMENTATION / PROTOTYPE")

left_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(6), Inches(5))
tf = left_box.text_frame
p = tf.add_paragraph()
p.text = "WORKING SOFTWARE PROTOTYPE"
p.font.bold = True
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(22, 163, 74)

bullets = [
    "Dashboard UI",
    "Sensor simulation",
    "Health visualization",
    "Condition indication",
    "Trend visualization",
    "System pipeline",
    "Dataset integration plan"
]
for b in bullets:
    p = tf.add_paragraph()
    p.text = "✓ " + b
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(17, 24, 39)

p2 = tf.add_paragraph()
p2.text = "\nLIVE PROTOTYPE:\nhttps://bridgeguard-ai-alpha.vercel.app/"
p2.font.bold = True
p2.font.size = Pt(14)
p2.font.color.rgb = RGBColor(37, 99, 235)

p3 = tf.add_paragraph()
p3.text = "\nCURRENT STAGE:\nSoftware prototype + simulation"
p3.font.bold = True
p3.font.size = Pt(14)
p3.font.color.rgb = RGBColor(22, 163, 74)

p4 = tf.add_paragraph()
p4.text = "\nIN PROGRESS:\n→ Real dataset preprocessing\n→ Health-relative feature extraction\n→ ML comparison & SHAP integration\n→ Hardware integration"
p4.font.bold = True
p4.font.size = Pt(14)
p4.font.color.rgb = RGBColor(245, 158, 11)

dash_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.5), Inches(1.5), Inches(6), Inches(4.5))
dash_box.fill.solid()
dash_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
dash_box.line.color.rgb = RGBColor(203, 213, 225)
tf = dash_box.text_frame
p = tf.add_paragraph()
p.text = "[ BRIDGEGUARD ML WEB DASHBOARD ]\n(Insert Dashboard Screenshot Here)"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = RGBColor(100, 116, 139)
p.alignment = PP_ALIGN.CENTER

add_speaker_notes(slide, "For our initial implementation, we have successfully developed the working software prototype and simulation dashboard. Current progress includes the UI and simulated health visualization. We are actively progressing on real dataset preprocessing, health-relative feature extraction, ML comparison, and hardware integration.")

# ---------------------------------------------------------
# SLIDE 15 — EXPECTED OUTPUT + FUTURE WORK
# ---------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)
set_footer(slide)
add_title(slide, "EXPECTED OUTPUT + FUTURE WORK")

lb = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(5.5), Inches(5))
tf = lb.text_frame
p = tf.add_paragraph()
p.text = "EXPECTED SYSTEM OUTPUT"
p.font.bold = True
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(37, 99, 235)

outs = [
    "Structural condition / anomaly indication",
    "Health-related indicators",
    "Feature trends",
    "Model confidence / anomaly score (where supported)",
    "SHAP feature contribution",
    "Maintenance-oriented alert"
]
for o in outs:
    p = tf.add_paragraph()
    p.text = "• " + o
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(17, 24, 39)
    p.space_after = Pt(5)

rb = slide.shapes.add_textbox(Inches(6.5), Inches(1.5), Inches(6), Inches(5))
tf = rb.text_frame
p = tf.add_paragraph()
p.text = "FUTURE WORK"
p.font.bold = True
p.font.size = Pt(18)
p.font.color.rgb = RGBColor(245, 158, 11)

fws = [
    "1. Preprocess real SHM datasets.",
    "2. Implement health-relative feature extraction.",
    "3. Establish healthy baseline.",
    "4. Compare ML algorithms.",
    "5. Train and validate selected model.",
    "6. Evaluate cross-condition / cross-bridge robustness.",
    "7. Integrate SHAP.",
    "8. Integrate ESP32 + sensors.",
    "9. Connect live sensor data to dashboard.",
    "10. Validate using miniature bridge experiments."
]
for f in fws:
    p = tf.add_paragraph()
    p.text = f
    p.font.size = Pt(14)
    p.font.color.rgb = RGBColor(17, 24, 39)
    p.space_after = Pt(2)

add_speaker_notes(slide, "Moving forward, our expected output includes anomaly indication and SHAP feature contributions. Our future work strictly follows the pipeline: preprocessing datasets, extracting relative features, evaluating ML robustness, integrating SHAP, and finally validating with physical ESP32 sensors on a miniature bridge.")

# ---------------------------------------------------------
# SLIDE 16 — REFERENCES
# ---------------------------------------------------------
slide = prs.slides.add_slide(blank_layout)
set_footer(slide)
add_title(slide, "REFERENCES")

ref_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12.33), Inches(5.5))
tf = ref_box.text_frame
refs = [
    "[1] Worden, K., & Manson, G. (2007). The application of machine learning to structural health monitoring.",
    "[2] Farrar, C. R., & Worden, K. (2012). Structural Health Monitoring: A Machine Learning Perspective.",
    "[3] Avci, O., et al. (2021). A review of vibration-based damage detection in civil structures...",
    "[4] Bao, Y., et al. (2019). Computer vision and machine learning for structural health monitoring.",
    "[5] Azimi, M., et al. (2020). Data-driven structural health monitoring and damage detection...",
    "[6] Entezami, A., et al. (2022). Structural health monitoring by a novel unsupervised machine learning method...",
    "[7] Neves, A. C., et al. (2017). Structural health monitoring of bridges: a model-free anomaly detection approach.",
    "[8] Malekzadeh, M., et al. (2015). Machine learning for structural health monitoring of bridges.",
    "[9] Sun, L., et al. (2020). A review of deep learning for structural health monitoring.",
    "[10] Figueiredo, E., et al. (2011). Machine learning algorithms for damage detection under operational variability.",
    "[11] Oh, B. K., et al. (2020). Explainable machine learning approach for structural health monitoring.",
    "[12] Toh, G., & Park, J. (2020). Review of Machine Learning Applications in Bridge Health Monitoring.",
    "[13] Vamvoudakis-Stefanou, K., et al. (2023). Unsupervised damage detection for bridge populations.",
    "[14] Sony, S., et al. (2021). A review of next-generation smart sensing technology in SHM.",
    "[15] Huang, Y., et al. (2019). A machine learning approach to anomaly detection in bridge monitoring data.",
    "[Dataset] Vänersborg Bridge SHM Dataset, Sweden. Zenodo. DOI: 10.5281/zenodo.8300495"
]

for r in refs:
    p = tf.add_paragraph()
    p.text = r
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(17, 24, 39)
    p.space_after = Pt(2)

add_speaker_notes(slide, "Here are our references, including 15 key academic papers covering ML for SHM, operational variability, and explainability, as well as the real-world Vänersborg Bridge SHM Dataset from Zenodo.")

prs.save("BridgeGuard_ML_Review_II.pptx")
print("Presentation generated successfully.")
